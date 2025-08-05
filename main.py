import sys
import json
import csv
from typing import List, Dict, Union
from src.processing import filter_by_state, sort_by_date
from src.utils import load_transactions, process_bank_search
from src.file_operations import read_csv, read_excel
from src.masks import get_mask_account, get_mask_card_number


def filter_rub_transactions(transactions: List[Dict]) -> List[Dict]:
    """
    Фильтрует транзакции, оставляя только рублёвые.
    Поддерживает все форматы (CSV, JSON, XLSX).
    """
    rub_transactions = []
    rub_aliases = {'RUB', 'РУБ', 'RUR', 'RU', '643'}  # 643 - цифровой код RUB

    for tx in transactions:
        try:
            # Проверяем все возможные места хранения валюты
            currency = (
                str(tx.get('operationAmount', {}).get('currency', {}).get('code', '')) or  # Для JSON
                str(tx.get('currency_code', '')) or  # Для XLSX/CSV
                str(tx.get('currency', '')) or  # Для старых форматов
                'RUB'
            ).upper().strip()

            if currency in rub_aliases:
                rub_transactions.append(tx)
        except:
            continue

    return rub_transactions


def safe_convert(value: Union[str, float, int]) -> float:
    """Безопасное преобразование в число"""
    if isinstance(value, (float, int)):
        return float(value)
    try:
        # Удаляем пробелы и заменяем запятые на точки
        cleaned = str(value).replace(' ', '').replace(',', '.')
        return float(cleaned)
    except (ValueError, TypeError):
        return 0.0


def print_transaction(transaction: Dict) -> None:
    """
    Улучшенная версия вывода транзакции с полной поддержкой JSON-структуры
    """
    try:
        # Проверка на пустую транзакцию
        if not transaction:
            print("\nПустая транзакция")
            return

        # Обработка даты
        raw_date = str(transaction.get('date', '')).strip()
        date = raw_date.split('T')[0] if 'T' in raw_date else raw_date
        date = date or 'Дата неизвестна'

        # Описание операции
        description = str(transaction.get('description', '')).strip()
        description = description or 'Без описания'

        print(f"\n{date} {description}")

        # Обработка отправителя/получателя (без изменений)
        for direction in ['from', 'to']:
            if direction not in transaction:
                continue

            value = str(transaction[direction]).strip()
            if not value:
                continue

            try:
                is_card = any(
                    word in value.lower()
                    for word in ['карта', 'card', 'visa', 'mastercard', 'discover', 'maestro']
                )

                parts = value.rsplit(' ', 1)
                if len(parts) == 2:
                    name, number = parts
                    masked = f"{name} {get_mask_card_number(number)}" if is_card else f"{name} {get_mask_account(number)}"
                else:
                    masked = value

                print(masked, end=' -> ' if direction == 'from' else '\n')
            except Exception as e:
                print(f"\nОшибка маскировки {direction}: {str(e)}")
                print(f"Исходное значение: {value}")

        # Получение суммы и валюты из JSON-структуры
        operation_amount = transaction.get('operationAmount', {})

        # Сумма может быть в operationAmount.amount или в корне транзакции
        amount = operation_amount.get('amount', transaction.get('amount'))

        # Валюта может быть в operationAmount.currency.code или в корне
        currency = operation_amount.get('currency', {}).get('code',
                                                            transaction.get('currency_code',
                                                                            transaction.get('currency', 'RUB'))).upper()

        # Вывод суммы
        if amount is not None:
            try:
                if isinstance(amount, str):
                    # Удаляем все нечисловые символы, кроме точки и запятой
                    amount_clean = ''.join(c for c in amount if c.isdigit() or c in '.,')
                    # Заменяем запятую на точку для корректного преобразования
                    amount_clean = float(amount_clean.replace(',', '.'))
                else:
                    amount_clean = float(amount)

                print(f"Сумма: {amount_clean:.2f} {currency}")
            except (ValueError, TypeError):
                print(f"Сумма: {amount} {currency} (некорректный формат)")
        else:
            print("Сумма: не указана")

    except Exception as e:
        print(f"\nКритическая ошибка при выводе транзакции: {str(e)}")
        print("Сырые данные:", transaction)



def load_json(filepath: str) -> List[Dict]:
    """Загрузка JSON файла с улучшенной обработкой суммы"""
    with open(filepath, 'r', encoding='utf-8') as f:
        data = json.load(f)

    processed_transactions = []

    # Обрабатываем как список, так и одиночную транзакцию
    transactions = data if isinstance(data, list) else [data]

    for tx in transactions:
        # Стандартизируем структуру
        processed = {
            'id': str(tx.get('id', '')),
            'state': str(tx.get('state', '')).upper(),
            'date': str(tx.get('date', '')),
            'description': str(tx.get('description', '')),
            'amount': tx.get('operationAmount', {}).get('amount', '0'),
            'currency': tx.get('operationAmount', {}).get('currency', {}).get('code', 'RUB'),
            'from': tx.get('from', ''),
            'to': tx.get('to', '')
        }
        processed_transactions.append(processed)

    return processed_transactions


def load_csv(filepath: str) -> List[Dict]:
    """Загрузка CSV файла с конкретным форматом"""
    transactions = []

    with open(filepath, 'r', encoding='windows-1251') as f:
        # Читаем файл с разделителем точка с запятой
        reader = csv.DictReader(f, delimiter=';')

        for row in reader:
            try:
                # Основные поля
                transaction = {
                    'id': row['id'].strip(),
                    'state': row['state'].strip().upper(),  # Приводим к верхнему регистру
                    'date': row['date'].strip().split('T')[0],  # Берем только дату без времени
                    'amount': float(row['amount'].replace(' ', '').replace(',', '.')),
                    'currency': row['currency_code'].strip().upper(),  # Используем currency_code как основной
                    'currency_name': row['currency_name'].strip(),
                    'currency_code': row['currency_code'].strip().upper(),
                    'from': row.get('from', '').strip(),  # Используем get() так как поле может отсутствовать
                    'to': row['to'].strip(),
                    'description': row['description'].strip()
                }
                transactions.append(transaction)

            except (KeyError, ValueError) as e:
                print(f"Ошибка обработки строки: {row}. Пропускаем. Ошибка: {str(e)}")
                continue

    return transactions


def load_xlsx(filepath: str) -> List[Dict]:
    """Загрузка XLSX файла с правильным извлечением валюты"""
    import openpyxl
    transactions = []

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # Получаем заголовки из первой строки
    headers = [str(cell.value).lower().strip() for cell in ws[1]]

    for row in ws.iter_rows(min_row=2):
        transaction = {}
        for header, cell in zip(headers, row):
            value = cell.value
            transaction[header] = str(value) if value is not None else ''

        # Стандартизируем формат валюты
        currency = (
            transaction.get('currency_code') or
            transaction.get('currency') or
            'RUB'
        ).upper()

        formatted = {
            'date': transaction.get('date', ''),
            'description': transaction.get('description', ''),
            'amount': transaction.get('amount', '0'),
            'currency': currency,  # Используем нормализованную валюту
            'currency_code': currency,  # Дублируем для совместимости
            'from': transaction.get('from', ''),
            'to': transaction.get('to', '')
        }
        transactions.append(formatted)

    return transactions



def process_file(filepath: str) -> None:
    """Обработка файла любого формата"""
    if filepath.endswith('.json'):
        transactions = load_json(filepath)
    elif filepath.endswith('.csv'):
        transactions = load_csv(filepath)
    elif filepath.endswith('.xlsx'):
        transactions = load_xlsx(filepath)
    else:
        raise ValueError("Неподдерживаемый формат файла")

    for idx, tx in enumerate(transactions, 1):
        print(f"\nТранзакция #{idx}:")
        print_transaction(tx)
        print("-" * 40)


def display_transactions(transactions: List[Dict]) -> None:
    """Выводит список транзакций с заголовком"""
    if not transactions:
        print("\nНе найдено ни одной транзакции, подходящей под ваши условия фильтрации")
        return

    print("\nРаспечатываю итоговый список транзакций...")
    print(f"Всего банковских операций в выборке: {len(transactions)}\n")
    for transaction in transactions:
        print_transaction(transaction)
        print()


def file_tip_transactions(filename: str, file_type: str) -> List[Dict]:
    """Загружает транзакции в зависимости от типа файла"""
    try:
        if file_type == '1':
            return load_transactions(filename)
        elif file_type == '2':
            return read_csv(filename)
        elif file_type == '3':
            return read_excel(filename)
    except Exception as e:
        raise Exception(f"Ошибка загрузки файла: {str(e)}")


def get_valid_input(prompt: str, valid_options: list) -> str:
    """Получает корректный ввод от пользователя"""
    while True:
        user_input = input(prompt).strip().lower()
        if user_input in valid_options:
            return user_input
        print(f"Ошибка: допустимые варианты - {', '.join(valid_options)}")


def main():
    print("Привет! Добро пожаловать в программу работы с банковскими транзакциями.")

    # Шаг 1: Выбор типа файла
    print("\nВыберите необходимый пункт меню:")
    print("1. Получить информацию о транзакциях из JSON-файла")
    print("2. Получить информацию о транзакциях из CSV-файла")
    print("3. Получить информацию о транзакциях из XLSX-файла")

    file_type = get_valid_input("Ваш выбор (1-3): ", ['1', '2', '3'])
    filename = input("Введите путь к файлу: ")

    try:
        # Шаг 2: Загрузка данных
        transactions = file_tip_transactions(filename, file_type)
        file_types = {
            '1': 'JSON',
            '2': 'CSV',
            '3': 'XLSX'
        }
        print(f"\nДля обработки выбран {file_types[file_type]}-файл.")
    except FileNotFoundError:
        print("\nОшибка: указанный файл не найден", file=sys.stderr)
        return
    except Exception as e:
        print(f"\n{e}", file=sys.stderr)
        return

    # Шаг 3: Фильтрация по статусу
    valid_statuses = ['executed', 'canceled', 'pending']
    while True:
        status = input("\nВведите статус (EXECUTED/CANCELED/PENDING): ").strip().lower()
        if status in valid_statuses:
            filtered = filter_by_state(transactions, status.upper())
            print(f"\nОперации отфильтрованы по статусу \"{status.upper()}\"")
            break
        print(f"\nСтатус операции \"{status}\" недоступен.")

    # Шаг 4: Дополнительные фильтры
    if get_valid_input("\nОтсортировать по дате? (да/нет): ", ['да', 'нет']) == 'да':
        order = get_valid_input("По возрастанию или убыванию? (возрастанию/убыванию): ",
                                ['возрастанию', 'убыванию'])
        filtered = sort_by_date(filtered, reverse=(order == 'убыванию'))

    if get_valid_input("\nТолько рублевые транзакции? (да/нет): ", ['да', 'нет']) == 'да':
        filtered = filter_rub_transactions(filtered)

    if get_valid_input("\nФильтровать по слову в описании? (да/нет): ", ['да', 'нет']) == 'да':
        search_word = input("Введите слово для поиска: ").strip()
        filtered = process_bank_search(filtered, search_word)

    # Шаг 5: Вывод результатов
    print(display_transactions(filtered))

if __name__ == "__main__":
    main()
